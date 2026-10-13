"""Read-only PDF/XLSX reports from the approved Analytics V2 responses."""
from datetime import datetime, timezone
from decimal import Decimal
from html import escape
from io import BytesIO

from fastapi.responses import Response
from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from reportlab.lib import colors
from reportlab.lib.pagesizes import A4
from reportlab.lib.styles import ParagraphStyle, getSampleStyleSheet
from reportlab.lib.units import mm
from reportlab.platypus import Image, LongTable, Paragraph, SimpleDocTemplate, Spacer, Table, TableStyle

from app.core.official_identity import (
    ADMINISTRATION_SECRETARIAT, COLOR_NAVY, FLEET_DEPARTMENT, MUNICIPALITY_NAME,
    crest_path, ensure_pdf_fonts,
)
from app.services.analytics_service import ANALYTICS_EXPORT_NO_CACHE_HEADERS, _analytics_xlsx_value, _institutional_timestamp


PROFILE_NAMES = {
    'executive': 'Visão geral', 'costs': 'Custos', 'fuel': 'Combustível',
    'maintenance': 'Manutenção', 'utilization': 'Utilização', 'drivers': 'Condutores',
}
QUALITY_LABELS = {
    'records': 'Registros',
    'operational_cost_missing_or_invalid': 'Custos sem valor válido',
    'events_without_driver': 'Eventos sem condutor vinculado',
    'distance_status': 'Cobertura da quilometragem',
}
QUALITY_VALUES = {
    'recorded': 'registrado',
    'partial': 'parcial',
    'unavailable': 'indisponível',
    'closed_possessions_partial_coverage': 'parcial: somente posses encerradas válidas',
}
METHODOLOGY_LABELS = {
    'cost_basis': 'Base de custo', 'claim_estimate': 'Sinistros estimados',
    'organization': 'Secretaria', 'vehicle_universe': 'Universo de veículos',
    'driver_risk': 'Indicador V1 de condutores', 'period': 'Período',
    'reference': 'Referência', 'mileage': 'Quilometragem',
    'persistence': 'Persistência', 'serialization': 'Serialização',
    'cost': 'Custo', 'claims': 'Sinistros', 'missing': 'Valores ausentes',
    'liters': 'Litros', 'price': 'Preço', 'category': 'Categoria',
    'anomalies': 'Anomalias', 'cohort': 'Recorte de intervenções',
    'duration': 'Duração', 'open': 'Abertas', 'recurrence': 'Reincidência',
    'roster': 'Cadastro atual', 'events': 'Eventos', 'distance': 'Distância',
    'last': 'Último evento', 'without': 'Sem evento', 'status': 'Status',
    'ranking': 'Ranking', 'scope': 'Escopo', 'possession': 'Posses',
    'incidents': 'Multas e sinistros', 'fuel': 'Abastecimentos',
    'license': 'CNH', 'coverage': 'Cobertura',
}
ANOMALY_LABELS = {'capacity': 'Capacidade do tanque', 'odometer': 'Hodômetro',
    'close': 'Abastecimentos próximos', 'consumption': 'Razão de intervalo km/L', 'price': 'Preço por litro'}
SOURCE_LABELS = {'possession': 'Posses', 'fine': 'Multas', 'claim': 'Sinistros',
    'fuel_supply': 'Abastecimentos', 'anomaly_rule': 'Regras de anomalia'}


def _value(value, unit=''):
    if value is None:
        return 'Não disponível'
    if isinstance(value, bool):
        return 'Sim' if value else 'Não'
    if isinstance(value, (Decimal, float)):
        value = f'{value:.2f}'
    elif isinstance(value, str):
        value = QUALITY_VALUES.get(value, value)
    return f'{value}{" " + unit if unit else ""}'


def _amount(amount, unit='R$'):
    if amount.value is not None:
        return _value(amount.value, unit)
    return f'Indisponível; subtotal conhecido {_value(amount.known_value, unit)} ({amount.missing_or_invalid} sem valor válido)'


def _section(title, rows):
    return title, [(str(label), _value(value), str(note or '')) for label, value, note in rows]


def report_sections(profile, data):
    """Curated fields only; never serialize entire domain responses into a report."""
    sections = []
    if profile == 'executive':
        sections.append(_section('Indicadores', [
            (item.label, _value(item.value, item.unit),
             f'{item.calculation.formula} Fonte: {", ".join(item.calculation.sources)}. '
             f'Período anterior: {_value(item.comparison.value, item.unit)}; '
             f'variação: {_value(item.comparison.delta, item.unit)} '
             f'({_value(item.comparison.delta_percent, "%")}). Qualidade: {_value(item.quality)}.')
            for item in data.kpis
        ]))
        sections.append(_section('Custos registrados e estimativa', [
            ('Custo operacional registrado', _amount(data.current.operational_cost), 'Combustível + manutenção + multas registradas.'),
            ('Sinistros estimados', _amount(data.current.claim_estimate), 'Estimativa separada do custo operacional registrado.'),
        ]))
        sections.append(_section('Evolução mensal', [
            (item.month, _amount(item.totals.operational_cost),
             f'Sinistros estimados, separados: {_amount(item.totals.claim_estimate)}')
            for item in data.monthly
        ]))
        sections.append(_section('Qualidade e cobertura', [
            (QUALITY_LABELS.get(key, key.replace('_', ' ').capitalize()), value, '') for key, value in data.quality.items()
        ]))
    elif profile == 'costs':
        sections.append(_section('Composição', [
            ('Custo operacional registrado', _amount(data.totals.operational_cost), 'Combustível + manutenção + multas registradas.'),
            ('Combustível registrado', _amount(data.totals.fuel), ''),
            ('Manutenção registrada', _amount(data.totals.maintenance), ''),
            ('Multas registradas', _amount(data.totals.fines), 'Todos os status; não significa pagamento.'),
            ('Sinistros estimados', _amount(data.totals.claim_estimate), 'Separados do custo realizado.'),
            ('Custo operacional por km', _value(data.cost_per_km, 'R$/km'), data.methodology.get('mileage', '')),
            ('Custo do conjunto medido', _amount(data.measured_cost), 'Numerador do custo/km.'),
            ('Km do conjunto medido', _value(data.measured_distance_km, 'km'), f'{data.measured_vehicles} veículos; denominador do custo/km.'),
        ]))
        sections.append(_section('Veículos por custo operacional registrado', [
            (item.plate, _amount(item.totals.operational_cost),
             f'Km válido: {_value(item.distance_km, "km")}; custo/km: {_value(item.cost_per_km, "R$/km")}')
            for item in data.vehicles
        ]))
        sections.append(_section('Secretarias por responsabilidade do registro', [
            (item.name, _amount(item.totals.operational_cost),
             f'{item.vehicle_count} veículo(s); sinistros estimados separados: {_amount(item.totals.claim_estimate)}')
            for item in data.organizations
        ]))
        sections.append(_section('Evolução mensal', [
            (item.month, _amount(item.totals.operational_cost),
             f'Sinistros estimados separados: {_amount(item.totals.claim_estimate)}')
            for item in data.monthly
        ]))
    elif profile == 'fuel':
        sections.append(_section('Abastecimentos', [
            ('Abastecimentos registrados', data.totals.records, ''),
            ('Litros abastecidos', _amount(data.totals.liters, 'L'), 'Não representa combustível consumido.'),
            ('Valor registrado', _amount(data.totals.cost), ''),
            ('Preço por litro ponderado', _value(data.totals.price_per_liter, 'R$/L'), 'Valor / litros dos registros com ambos válidos.'),
            ('Km por litro abastecido, conjunto medido', _value(data.measured_km_per_liter_supplied, 'km/L'), data.methodology.get('mileage', '')),
            ('Litros abastecidos por 100 km, conjunto medido', _value(data.measured_liters_per_100km_supplied, 'L/100km'), ''),
            ('Km de posses válidas dos veículos abastecidos', _value(data.measured_distance_km, 'km'), f'{data.measured_vehicles} veículos medidos.'),
        ]))
        sections.append(_section('Acionamentos das regras de anomalia', [
            (ANOMALY_LABELS.get(kind, kind), count, data.methodology.get('anomalies', '')) for kind, count in data.anomaly_counts.items()
        ]))
        sections.append(_section('Evolução mensal', [
            (item.month, _amount(item.totals.cost), f'Litros abastecidos: {_amount(item.totals.liters, "L")}')
            for item in data.monthly
        ]))
        sections.append(_section('Postos registrados', [
            (item.name, _amount(item.totals.cost),
             f'{item.totals.records} abastecimento(s); {_amount(item.totals.liters, "L")}; preço ponderado: {_value(item.totals.price_per_liter, "R$/L")}; {item.anomaly_count} acionamento(s) de regra')
            for item in data.stations
        ]))
        sections.append(_section('Veículos', [
            (item.plate, _amount(item.totals.cost),
             f'{item.totals.records} abastecimento(s); {_amount(item.totals.liters, "L")}; km válido: {_value(item.distance_km, "km")}; '
             f'razão: {_value(item.km_per_liter_supplied, "km/L")}; custo/km: {_value(item.fuel_cost_per_km, "R$/km")}; '
             f'categoria: {_value(item.category_liters_per_100km, "L/100km")} ({item.category_sample_vehicles} pares); {item.anomaly_count} acionamento(s) de regra')
            for item in data.vehicles
        ]))
        if data.can_view_records:
            sections.append(_section('Evidências das anomalias', [
                (f'{item.plate} · {item.label}', item.observed,
                 f'{item.supplied_at.isoformat()} · Referência: {item.reference}; amostra: {item.sample_size}; regra: {item.rule}; {item.limitation}')
                for item in data.anomalies
            ]))
    elif profile == 'maintenance':
        sections.append(_section('Manutenções', [
            ('Custo registrado', _amount(data.cost), ''),
            ('Intervenções iniciadas', data.interventions, ''),
            ('Abertas no estado atual', data.open_interventions, 'Não reconstitui indisponibilidade histórica.'),
            ('Duração média válida', _value(data.average_duration_hours, 'h'), f'{data.valid_duration_count} intervalos válidos.'),
            ('Veículos com duas ou mais intervenções', data.repeated_vehicles, 'Reincidência quantitativa, sem atribuição de causa.'),
            ('Custo de manutenção por km', _value(data.cost_per_km, 'R$/km'), data.methodology.get('mileage', '')),
            ('Intervalos encerrados inválidos', data.invalid_duration_count, 'Excluídos da duração média.'),
            ('Custo do conjunto medido', _amount(data.measured_cost), 'Numerador do custo/km.'),
            ('Km do conjunto medido', _value(data.measured_distance_km, 'km'), f'{data.measured_vehicles} veículos medidos.'),
        ]))
        sections.append(_section('Veículos', [
            (item.plate, _amount(item.cost), f'{item.interventions} intervenções; {item.open_interventions} abertas; '
             f'duração média: {_value(item.average_duration_hours, "h")} ({item.valid_duration_count} intervalos); '
             f'km válido: {_value(item.distance_km, "km")}; custo/km: {_value(item.cost_per_km, "R$/km")}')
            for item in data.vehicles
        ]))
    elif profile == 'utilization':
        sections.append(_section('Posse, hodômetro e cadastro', [
            ('Veículos no cadastro atual', data.roster_vehicles, ''),
            ('Posses iniciadas', data.started, ''),
            ('Posses encerradas', data.ended, ''),
            ('Km de posses encerradas válidas', _value(data.distance_km, 'km'), f'{data.valid_km_count} posses válidas.'),
            ('Km por posse encerrada válida', _value(data.km_per_valid_closed_possession, 'km'), ''),
            ('Sem evento de posse no recorte', data.without_possession_event, 'Posse não comprova deslocamento contínuo.'),
            ('Veículos com manutenção aberta agora', data.vehicles_with_open_maintenance, 'Estado atual, não disponibilidade histórica.'),
            ('Duração registrada de posses válidas', _value(data.duration_hours, 'h'), f'{data.valid_duration_count} intervalos; não são horas de motor ligado.'),
            ('Manutenções abertas agora', data.open_maintenance_records, 'Estado atual, sem reconstrução histórica.'),
        ]))
        sections.append(_section('Veículos', [
            (item.plate, _value(item.distance_km, 'km'),
             f'{item.started} posses iniciadas; {item.ended} encerradas; status atual: {item.status}; '
             f'duração registrada: {_value(item.duration_hours, "h")}; último evento: {_value(item.last_possession_event_at)}; '
             f'dias até o fim do recorte: {_value(item.days_since_last_event)}; {item.open_maintenance_records} manutenções abertas agora')
            for item in data.vehicles
        ]))
        sections.append(_section('Status cadastral atual', [
            (key, count, 'Situação atual; não reconstitui disponibilidade no período.')
            for key, count in data.status_counts.items()
        ]))
        sections.append(_section('Secretarias operadoras atuais', [
            (item.name, item.vehicles,
             f'{item.started} posses iniciadas; {item.ended} encerradas; {_value(item.distance_km, "km")} válidos')
            for item in data.organizations
        ]))
    elif profile == 'drivers':
        sections.append(_section('Eventos vinculados explicitamente', [
            ('Condutores no recorte', data.totals.get('drivers'), ''),
            ('Posses', data.totals.get('possessions'), ''),
            ('Km válido de posses encerradas', _value(data.totals.get('distance_km'), 'km'), ''),
            ('Multas', data.totals.get('fines'), ''),
            ('Sinistros', data.totals.get('claims'), ''),
            ('Abastecimentos', data.totals.get('fuel_supplies'), ''),
            ('Acionamentos de regra de anomalia', data.totals.get('anomaly_rules'), 'Não indicam culpa do condutor.'),
        ]))
        sections.append(_section('Condutores', [
            (item.name, _value(item.distance_km, 'km'),
             f'{item.possessions} posses; {item.valid_km_possessions} com km válido; {item.vehicles_with_possession} veículos; '
             f'{item.fines} multas; {item.claims} sinistros; {item.fuel_supplies} abastecimentos; {item.anomaly_rules} regras')
            for item in data.drivers
        ]))
        sections.append(_section('Registros sem condutor vinculado', [
            (SOURCE_LABELS.get(key, key.replace('_', ' ').capitalize()), count, 'Sem atribuição por nome.')
            for key, count in data.unlinked_records.items()
        ]))
        if data.license is not None:
            sections.append(_section('CNH — indicador administrativo atual', [
                ('Vencidas', data.license.expired, ''), ('A vencer em 30 dias', data.license.expiring, ''),
                ('Sem data cadastrada', data.license.missing, ''),
                ('Condutores ativos no cadastro atual', data.license.active_drivers, f'Situação administrativa em {data.license.as_of.isoformat()}; independente do período histórico.'),
            ]))
    methodology = getattr(data, 'methodology', {})
    sections.append(_section('Fontes, critérios e limitações', [
        (METHODOLOGY_LABELS.get(key, key.replace('_', ' ').capitalize()), '', text)
        for key, text in methodology.items()
    ]))
    return sections


def _header_lines(profile, filters, issued_at):
    return [
        f'Relatório de {PROFILE_NAMES[profile]}',
        f'Período: {filters.date_from.isoformat()} a {filters.date_to.isoformat()} · dias civis encerrados (America/Bahia)',
        f'Secretaria: {filters.organization or "todas no escopo autorizado"} · tipo: {filters.vehicle_type.value if filters.vehicle_type else "todos"} · veículo: {filters.vehicle_id or "todos"}',
        f'Emissão: {_institutional_timestamp(issued_at)}',
    ]


def _xlsx(profile, filters, sections, issued_at):
    workbook = Workbook()
    workbook.properties.creator = MUNICIPALITY_NAME
    workbook.properties.title = f'Relatório de {PROFILE_NAMES[profile]}'
    sheet = workbook.active
    sheet.title = 'Relatório'
    sheet.sheet_view.showGridLines = False
    row = 1
    for line in (MUNICIPALITY_NAME, f'{ADMINISTRATION_SECRETARIAT} · {FLEET_DEPARTMENT}', *_header_lines(profile, filters, issued_at)):
        cell = sheet.cell(row=row, column=1, value=_analytics_xlsx_value(line))
        cell.font = Font(bold=row in (1, 3), color=COLOR_NAVY.lstrip('#'))
        sheet.merge_cells(start_row=row, start_column=1, end_row=row, end_column=3)
        row += 1
    row += 1
    for title, entries in sections:
        sheet.cell(row=row, column=1, value=_analytics_xlsx_value(title)).font = Font(bold=True, color='FFFFFF')
        for column in range(1, 4):
            sheet.cell(row=row, column=column).fill = PatternFill('solid', fgColor=COLOR_NAVY.lstrip('#'))
        row += 1
        for column, label in enumerate(('Indicador / entidade', 'Valor', 'Fonte / contexto'), 1):
            sheet.cell(row=row, column=column, value=label).font = Font(bold=True)
        row += 1
        for entry in entries or [('Sem registros no recorte', '', '')]:
            for column, value in enumerate(entry, 1):
                cell = sheet.cell(row=row, column=column, value=_analytics_xlsx_value(value))
                cell.alignment = Alignment(vertical='top', wrap_text=True)
            row += 1
        row += 1
    for column, width in {'A': 38, 'B': 43, 'C': 100}.items():
        sheet.column_dimensions[column].width = width
    sheet.freeze_panes = 'A8'
    sheet.page_setup.orientation = 'landscape'
    output = BytesIO()
    workbook.save(output)
    return output.getvalue()


def _pdf(profile, filters, sections, issued_at):
    output = BytesIO()
    regular, bold = ensure_pdf_fonts()
    styles = getSampleStyleSheet()
    body = ParagraphStyle('ReportBody', parent=styles['BodyText'], fontName=regular, fontSize=8, leading=11)
    heading = ParagraphStyle('ReportHeading', parent=body, fontName=bold, fontSize=11, leading=14,
                             textColor=colors.HexColor(COLOR_NAVY), spaceBefore=5 * mm, keepWithNext=True)
    small = ParagraphStyle('ReportSmall', parent=body, fontSize=7, leading=9)
    table_header = ParagraphStyle('ReportTableHeader', parent=small, fontName=bold, textColor=colors.white)
    title = ParagraphStyle('ReportTitle', parent=heading, fontSize=15, leading=18)
    document = SimpleDocTemplate(output, pagesize=A4, leftMargin=14 * mm, rightMargin=14 * mm,
                                 topMargin=12 * mm, bottomMargin=15 * mm,
                                 title=f'Relatório de {PROFILE_NAMES[profile]}', author=MUNICIPALITY_NAME)
    story = [Image(str(crest_path()), width=10 * mm, height=12.6 * mm),
             Paragraph(escape(MUNICIPALITY_NAME), heading),
             Paragraph(escape(f'{ADMINISTRATION_SECRETARIAT} · {FLEET_DEPARTMENT}'), small)]
    for index, line in enumerate(_header_lines(profile, filters, issued_at)):
        story.append(Paragraph(escape(line), title if index == 0 else small))
    for section_title, entries in sections:
        story.append(Spacer(1, 2 * mm))
        story.append(Paragraph(escape(section_title), heading))
        rows = [[Paragraph(escape(text), table_header) for text in ('Indicador / entidade', 'Valor', 'Fonte / contexto')]]
        rows.extend([[Paragraph(escape(str(value)), small) for value in entry]
                     for entry in (entries or [('Sem registros no recorte', '', '')])])
        table = LongTable(rows, colWidths=[48 * mm, 40 * mm, 94 * mm], repeatRows=1, hAlign='LEFT')
        table.setStyle(TableStyle([
            ('BACKGROUND', (0, 0), (-1, 0), colors.HexColor(COLOR_NAVY)),
            ('TEXTCOLOR', (0, 0), (-1, 0), colors.white),
            ('GRID', (0, 0), (-1, -1), 0.3, colors.HexColor('#D7E0ED')),
            ('VALIGN', (0, 0), (-1, -1), 'TOP'),
            ('ROWBACKGROUNDS', (0, 1), (-1, -1), [colors.white, colors.HexColor('#F8FAFD')]),
        ]))
        story.append(table)
    document.build(story)
    return output.getvalue()


def export_report(profile, export_format, filters, data):
    issued_at = datetime.now(timezone.utc)
    sections = report_sections(profile, data)
    content = _xlsx(profile, filters, sections, issued_at) if export_format == 'xlsx' else _pdf(profile, filters, sections, issued_at)
    return Response(content=content, media_type=(
        'application/vnd.openxmlformats-officedocument.spreadsheetml.sheet' if export_format == 'xlsx' else 'application/pdf'),
        headers={**ANALYTICS_EXPORT_NO_CACHE_HEADERS,
                 'Content-Disposition': f'attachment; filename="analytics-v2-{profile}-{filters.date_from}-{filters.date_to}.{export_format}"'})

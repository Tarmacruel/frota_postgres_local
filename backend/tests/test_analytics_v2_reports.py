from datetime import date
from decimal import Decimal
from io import BytesIO
from types import SimpleNamespace
from unittest.mock import AsyncMock
from uuid import uuid4

import pytest
from fastapi import FastAPI
from httpx import ASGITransport, AsyncClient
from openpyxl import load_workbook

from app.api.deps import get_current_user_ready
from app.api.routes import analytics_v2 as routes
from app.db.session import get_db_session
from app.models.user import UserRole
from app.schemas.analytics_v2 import AnalyticsV2Amount, AnalyticsV2Filter
from app.services.analytics_v2_reports import export_report, report_sections


def report_filters(**overrides):
    return AnalyticsV2Filter(date_from=date(2026, 9, 1), date_to=date(2026, 9, 30), **overrides)


def amount(value, *, missing=0):
    return AnalyticsV2Amount(value=value if not missing else None, known_value=value or 0,
                             records=1, missing_or_invalid=missing)


def cost_data(**overrides):
    totals = SimpleNamespace(operational_cost=amount(Decimal('100.00')), fuel=amount(Decimal('70.00')),
        maintenance=amount(Decimal('20.00')), fines=amount(Decimal('10.00')),
        claim_estimate=amount(Decimal('999.00')))
    return SimpleNamespace(totals=totals, cost_per_km=Decimal('2.00'), measured_cost=totals.operational_cost,
        measured_distance_km=Decimal('50'), measured_vehicles=1, organizations=[], monthly=[], vehicles=[
        SimpleNamespace(plate='=HYPERLINK("bad")', totals=totals, distance_km=Decimal('50'), cost_per_km=Decimal('2'))],
        methodology={'cost': '+SUM(A1:A2)', 'mileage': 'Posses encerradas válidas.'}, **overrides)


def test_cost_report_separates_estimate_and_formula_safe_xlsx():
    data = cost_data()
    sections = report_sections('costs', data)
    composition = sections[0][1]
    assert composition[0][0] == 'Custo operacional registrado'
    assert '100.00' in composition[0][1]
    assert composition[4][0] == 'Sinistros estimados'
    assert '999.00' in composition[4][1]
    response = export_report('costs', 'xlsx', report_filters(), data)
    assert response.status_code == 200 and response.body.startswith(b'PK')
    assert response.headers['content-disposition'] == 'attachment; filename="analytics-v2-costs-2026-09-01-2026-09-30.xlsx"'
    assert response.headers['cache-control'].startswith('private, no-store')
    workbook = load_workbook(BytesIO(response.body), data_only=False)
    cells = [cell for row in workbook.active for cell in row]
    values = [cell.value for cell in cells]
    assert any('Sinistros estimados' == value for value in values)
    assert any('Período: 2026-09-01 a 2026-09-30' in str(value) for value in values)
    assert "'=HYPERLINK(\"bad\")" in values
    assert "'+SUM(A1:A2)" in values
    assert all(cell.data_type != 'f' for cell in cells)
    workbook.close()


def test_executive_report_formats_ratios_and_coverage_for_reading():
    kpi = SimpleNamespace(label='Custo operacional registrado por km', value=Decimal('0.018656516747391062'),
        unit='BRL/km', quality='partial', calculation=SimpleNamespace(formula='custo / km', sources=['posses']),
        comparison=SimpleNamespace(value=Decimal('0.045635'), delta=Decimal('-0.026978'),
                                   delta_percent=Decimal('-59.118')))
    data = SimpleNamespace(kpis=[kpi], current=SimpleNamespace(operational_cost=amount(Decimal('100.00')),
        claim_estimate=amount(Decimal('20.00'))), monthly=[], quality={
        'distance_status': 'closed_possessions_partial_coverage'},
        methodology={'mileage': 'Somente posses encerradas válidas.'})
    sections = report_sections('executive', data)
    indicator = sections[0][1][0]
    assert indicator[1] == '0.02 BRL/km'
    assert '0.05 BRL/km' in indicator[2] and '-0.03 BRL/km' in indicator[2]
    assert '(-59.12 %)' in indicator[2] and 'Qualidade: parcial.' in indicator[2]
    assert sections[3][1][0] == ('Cobertura da quilometragem',
        'parcial: somente posses encerradas válidas', '')
    assert sections[4][1][0] == ('Quilometragem', '', 'Somente posses encerradas válidas.')


def test_pdf_contains_profile_and_private_headers():
    response = export_report('costs', 'pdf', report_filters(), cost_data())
    assert response.body.startswith(b'%PDF-')
    assert response.media_type == 'application/pdf'
    assert response.headers['x-content-type-options'] == 'nosniff'


@pytest.mark.parametrize('profile,data,label,value', [
    ('fuel', SimpleNamespace(totals=SimpleNamespace(records=1, liters=amount(Decimal('10')),
        cost=amount(Decimal('50')), price_per_liter=Decimal('5')), measured_km_per_liter_supplied=Decimal('10'),
        measured_liters_per_100km_supplied=Decimal('10'), measured_distance_km=Decimal('100'), measured_vehicles=1,
        anomaly_counts={}, monthly=[], stations=[], vehicles=[], can_view_records=False, methodology={}),
        'Litros abastecidos por 100 km, conjunto medido', '10.00 L/100km'),
    ('maintenance', SimpleNamespace(cost=amount(Decimal('100')), interventions=2, open_interventions=1,
        average_duration_hours=Decimal('6'), valid_duration_count=1, repeated_vehicles=1,
        cost_per_km=Decimal('2'), invalid_duration_count=0, measured_cost=amount(Decimal('100')),
        measured_distance_km=Decimal('50'), measured_vehicles=1, vehicles=[], methodology={}),
        'Km do conjunto medido', '50.00 km'),
    ('utilization', SimpleNamespace(roster_vehicles=2, started=1, ended=1, distance_km=Decimal('50'),
        valid_km_count=1, km_per_valid_closed_possession=Decimal('50'), without_possession_event=1,
        vehicles_with_open_maintenance=1, duration_hours=Decimal('12'), valid_duration_count=1,
        open_maintenance_records=2, vehicles=[], status_counts={}, organizations=[], methodology={}),
        'Duração registrada de posses válidas', '12.00 h'),
])
def test_reports_include_measured_coverage_and_registered_duration(profile, data, label, value):
    response = export_report(profile, 'xlsx', report_filters(), data)
    workbook = load_workbook(BytesIO(response.body), data_only=False)
    rows = list(workbook.active.values)
    assert any(row[0] == label and row[1] == value for row in rows)
    workbook.close()


def _app(*, role=UserRole.ADMIN, organization=None, analytics_allowed=True, drivers_view=True):
    app = FastAPI()
    app.include_router(routes.router)
    user = SimpleNamespace(id=uuid4(), role=role, organization_id=organization, permissions={
        'analytics': {'can_view': analytics_allowed}, 'drivers': {'can_view': drivers_view},
        'fuel_supplies': {'can_view': True}, 'possession': {'can_view': True},
        'fines': {'can_view': True}, 'claims': {'can_view': True},
    })
    entry = SimpleNamespace(can_view=analytics_allowed, can_create=False, can_edit=False, can_delete=False)
    db = SimpleNamespace(execute=AsyncMock(return_value=SimpleNamespace(scalar_one_or_none=lambda: entry)))
    app.dependency_overrides[get_current_user_ready] = lambda: user
    app.dependency_overrides[get_db_session] = lambda: db
    return app


@pytest.mark.asyncio
async def test_export_route_uses_effective_scope_and_only_supported_parameters(monkeypatch):
    organization = uuid4()
    captured = []

    async def get(self, filters):
        captured.append(filters)
        return cost_data()

    monkeypatch.setattr(routes.AnalyticsV2Costs, 'get', get)
    app = _app(role=UserRole.PRODUCAO, organization=organization)
    base = {'date_from': '2026-09-01', 'date_to': '2026-09-30', 'profile': 'costs', 'export_format': 'xlsx'}
    async with AsyncClient(transport=ASGITransport(app=app), base_url='http://test') as client:
        response = await client.get('/api/analytics/v2/export', params={**base, 'organization': str(uuid4())})
        ignored = await client.get('/api/analytics/v2/export', params={**base, 'include_charts': 'true'})
        invalid = await client.get('/api/analytics/v2/export', params={**base, 'profile': 'personalized'})
    assert response.status_code == 200, response.text
    assert captured[0].organization == organization
    assert ignored.status_code == 422
    assert invalid.status_code == 422
    assert len(captured) == 1


@pytest.mark.asyncio
async def test_export_permission_and_driver_source_permission(monkeypatch):
    async def summary(self, filters):
        raise AssertionError('service should not run')
    monkeypatch.setattr(routes.AnalyticsV2Drivers, 'get', summary)
    params = {'date_from': '2026-09-01', 'date_to': '2026-09-30', 'profile': 'drivers', 'export_format': 'pdf'}
    async with AsyncClient(transport=ASGITransport(app=_app(analytics_allowed=False)), base_url='http://test') as client:
        assert (await client.get('/api/analytics/v2/export', params=params)).status_code == 403
    async with AsyncClient(transport=ASGITransport(app=_app(drivers_view=False)), base_url='http://test') as client:
        assert (await client.get('/api/analytics/v2/export', params=params)).status_code == 403
